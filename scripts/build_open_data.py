# /// script
# dependencies = ["openpyxl", "pandas"]
# ///
"""Build and refresh the repository's open data.

Usage:
  uv run --script build_open_data.py init  <root> <bulletins_daily.csv> <parse.log> <pdf_dir>
  uv run --script build_open_data.py daily <root>

init   one-off: today's parsed archive -> data/pump/YYYY.csv, data/pump/missing_days.csv, data/areas.csv
daily  refinery prices and the ECB rate -> data/ (committed);
       EIA spot prices and EIA's Short-Term Energy Outlook -> build/ (not committed: EIA's
       prices come from Refinitiv, an LSEG business, so they are fetched, not republished)
"""
import csv
import datetime as dt
import pathlib
import re
import sys
import urllib.request

import openpyxl
import pandas as pd

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from areas import AREAS, BY_LABEL, FUELS, write_csv

UA = {"User-Agent": "greek-fuel-prices (https://github.com/despotak/greek-fuel-prices)"}
REFINERY_URL = "https://raw.githubusercontent.com/akosiaris/greek-oil-refinery-prices/main/fuels.csv"
ECB_URL = "https://data-api.ecb.europa.eu/service/data/EXR/D.USD.EUR.SP00.A?format=csvdata&startPeriod=2017-01-01"
EIA_XLS = "https://www.eia.gov/dnav/pet/hist_xls/{}d.xls"
EIA_SERIES = ["EER_EPD2DXL0_PF4_Y35NY_DPG", "EER_EPMRU_PF4_Y35NY_DPG", "EER_EPD2F_PF4_Y35NY_DPG",
              "EER_EPLLPA_PF4_Y44MB_DPG", "RBRTE", "RWTC"]
STEO_URL = "https://www.eia.gov/outlooks/steo/xls/STEO_m.xlsx"
STEO_SERIES = {"BREPUUS": ("brent", "Brent crude oil spot price", "USD/bbl"),
               "DSWHUUS_$": ("diesel", "US diesel fuel, refiner wholesale price", "USD/gal"),
               "MGWHUUS_$": ("gasoline", "US gasoline, refiner wholesale price", "USD/gal")}
PUMP_COLUMNS = ["date", "area_id", *FUELS]


def get(url, path):
    path = pathlib.Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=120) as resp, open(path, "wb") as fh:
        fh.write(resp.read())
    return path


def write_pump_years(root, rows):
    """rows: [date, area_id, *FUELS] -> data/pump/YYYY.csv, merged with what is there, sorted."""
    order = {a[0]: i for i, a in enumerate(AREAS)}
    by_year = {}
    for r in rows:
        by_year.setdefault(r[0][:4], []).append(r)
    for year, new in by_year.items():
        path = root / "data" / "pump" / f"{year}.csv"
        old = []
        if path.exists():
            with open(path, encoding="utf-8") as fh:
                old = [list(x.values()) for x in csv.DictReader(fh)]
        merged = {(r[0], r[1]): r for r in old}
        merged.update({(r[0], r[1]): r for r in new})
        out = sorted(merged.values(), key=lambda r: (r[0], order[r[1]]))
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh, lineterminator="\n")
            w.writerow(PUMP_COLUMNS)
            w.writerows(out)


def init(root, daily_csv, parse_log, pdf_dir):
    root = pathlib.Path(root)
    write_csv(root / "data" / "areas.csv")
    rows = []
    with open(daily_csv, encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            rows.append([r["date"], BY_LABEL[r["region"]], *(r[f] for f in FUELS)])
    # Days no parser can read (scanned images, a broken text layout), read by hand from the page
    # image and kept as their own source so a rebuild from the archive keeps them.
    hand = []
    transcribed = root / "data" / "pump" / "transcribed.csv"
    if transcribed.exists():
        with open(transcribed, encoding="utf-8") as fh:
            hand = [[r["date"], r["area_id"], *(r[f] for f in FUELS)] for r in csv.DictReader(fh)]
    write_pump_years(root, rows + hand)

    parsed = {r[0] for r in rows} | {r[0] for r in hand}
    fails = {}
    notes = [[d, ("read by hand from the bulletin's page image (see transcribed.csv); checked against a machine "
                  "reading of the page (OCR, or the PDF's own text layer for 2017-03-27), the neighbouring days, and "
                  "the national average rebuilt from the prefecture prices for unleaded 95, 100 and diesel")]
             for d in sorted({r[0] for r in hand})]
    for line in pathlib.Path(parse_log).read_text(encoding="utf-8").splitlines():
        if (n := re.match(r"NOTE (\d{4}-\d{2}-\d{2}): (.*)", line)) and n[1] in parsed:
            notes += [[n[1], part] for part in n[2].split("; ")]
            if moved := re.search(r"body says ([\d-]+), filename says ([\d-]+)", n[2]):
                fails.setdefault(moved[2], "file for this day is a mislabelled copy of " + moved[1])
        m = re.match(r"FAIL (\d{4}-\d{2}-\d{2})(?:\.pdf)?: (?:\w+Error )?(.*)", line)
        if m:
            msg = m[2]
            if "no prices in the text layer" in msg:
                reason = "bulletin published as a scanned image or a notice, no readable prices"
            elif "skipping duplicate" in msg:
                reason = "file for this day is a mislabelled copy of " + re.search(r"body says ([\d-]+)", msg)[1]
            else:
                reason = "bulletin could not be read: " + msg
            fails[m[1]] = reason
    have_pdf = {p.stem for p in pathlib.Path(pdf_dir).glob("*.pdf")}
    first, last = min(parsed), max(parsed)
    missing = []
    for d in pd.date_range(first, last, freq="D").strftime("%Y-%m-%d"):
        if d in parsed:
            continue
        missing.append([d, fails.get(d) or ("bulletin could not be read" if d in have_pdf else "no bulletin published on fuelprices.gr")])
    with open(root / "data" / "pump" / "missing_days.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["date", "reason"])
        w.writerows(missing)
    with open(root / "data" / "pump" / "day_notes.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["date", "note"])
        w.writerows(sorted(notes))
    print(f"pump: {len(rows)} rows over {len(parsed)} days; missing days: {len(missing)}", file=sys.stderr)


def refinery(root, tmp):
    raw = pd.read_csv(get(REFINERY_URL, tmp / "refinery_raw.csv"))
    raw["date"] = raw["date"].str[:10]
    today = dt.date.today().isoformat()
    raw = raw[raw["date"] <= today].drop_duplicates(["date", "fuel"], keep="last")
    out = pd.DataFrame({
        "date": raw["date"], "category": raw["category"], "product": raw["fuel"],
        "elpe": raw["elpePrice"], "motor_oil": raw["motoroilPrice"], "mean": raw["meanPrice"],
        "mean_incl_vat": raw["vatPrice"], "unit": raw["unit"], "basis": raw["notes"],
    }).sort_values(["date", "product"])
    path = root / "data" / "refinery" / "refinery_prices.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(path, index=False, lineterminator="\n")
    print(f"refinery: {len(out)} rows, {out['date'].min()} .. {out['date'].max()}", file=sys.stderr)


def ecb(root, tmp):
    raw = pd.read_csv(get(ECB_URL, tmp / "ecb.csv"), usecols=["TIME_PERIOD", "OBS_VALUE"])
    out = raw.rename(columns={"TIME_PERIOD": "date", "OBS_VALUE": "usd_per_eur"}).sort_values("date")
    path = root / "data" / "international" / "ecb_eurusd.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(path, index=False, lineterminator="\n")
    print(f"ecb: {len(out)} rows to {out['date'].max()}", file=sys.stderr)


def eia(root):
    for s in EIA_SERIES:
        get(EIA_XLS.format(s), root / "build" / "eia" / f"{s}.xls")
    path = get(STEO_URL, root / "build" / "STEO_m.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # Labels in the "Dates" sheet sit in varying columns and carry trailing spaces and dashes,
    # so each row is reduced to its non-empty cells and matched on the first one by prefix.
    dates = [cells for r in wb["Dates"].iter_rows(values_only=True) if (cells := [c for c in r if c is not None])]

    def row(prefix):
        return next(cells for cells in dates if str(cells[0]).strip().startswith(prefix))

    months = [m for m in row("AAAA_DATEX")[1:] if isinstance(m, int)]
    last_hist = int(row("Last Historical Month")[1])
    vintage = row("Forecast Month")[1]
    rows = []
    for r in wb["2tab"].iter_rows(values_only=True):
        if r and r[0] in STEO_SERIES:
            key, name, unit = STEO_SERIES[r[0]]
            for m, v in zip(months, r[2:]):
                if isinstance(v, (int, float)):
                    rows.append([f"{m // 100}-{m % 100:02d}", key, name, round(float(v), 4), unit,
                                 "history" if m <= last_hist else "forecast", vintage])
    with open(root / "build" / "steo_monthly.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["month", "series", "name", "value", "unit", "kind", "vintage"])
        w.writerows(rows)
    print(f"eia: {len(EIA_SERIES)} spot series; steo {vintage}: {len(rows)} rows, last history {last_hist}", file=sys.stderr)


def daily(root):
    root = pathlib.Path(root)
    tmp = root / "build" / "tmp"
    refinery(root, tmp)
    ecb(root, tmp)
    eia(root)


if __name__ == "__main__":
    cmd, *args = sys.argv[1:]
    {"init": init, "daily": daily}[cmd](*args)
